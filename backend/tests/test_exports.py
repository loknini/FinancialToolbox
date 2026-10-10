import json
from datetime import timedelta
from decimal import Decimal
from io import BytesIO
from uuid import uuid4

import pytest
from openpyxl import load_workbook
from pypdf import PdfReader
from sqlalchemy import func, select

from app.core.errors import BusinessError
from app.core.security import decode, now
from app.db.models import Export, ExportFile, ExportShare, ExportSnapshot, IdempotencyRecord, Job
from app.integrations.ocr import normalize_response
from app.jobs import export_tasks
from app.services.export_render import summaries
from tests.helpers import local_url
from tests.test_invoices import drain, payload, setup_batch


def create_tax(client, auth, amount="113.00"):
    result = client.post(
        "/api/v1/tax-calculations",
        headers={**auth, "Idempotency-Key": str(uuid4())},
        json={
            "gross_amount": amount,
            "tax_rate": "0.13",
            "invoice_type_code": "general",
            "config_version": "tax-config-v1",
        },
    )
    assert result.status_code == 201, result.text
    return result.json()["data"]


def create_export(
    client, auth, *, format="excel", source="tax_calculations", selection=None, **kwargs
):
    body = {
        "format": format,
        "source_type": source,
        "selection": selection or {"mode": "filter"},
        **kwargs,
    }
    headers = {**auth, "Idempotency-Key": str(uuid4())}
    result = client.post("/api/v1/exports", headers=headers, json=body)
    assert result.status_code == 202, result.text
    return result.json()["data"], headers, body


def finished(client, auth, worker, row):
    drain(worker)
    value = client.get("/api/v1/exports/" + row["id"], headers=auth).json()["data"]
    assert value["status"] == "succeeded", value
    link = client.get("/api/v1/exports/" + row["id"] + "/download-url", headers=auth)
    assert link.status_code == 200, link.text
    result = client.get(local_url(link.json()["data"]["url"]))
    assert result.status_code == 200
    return value, result.content


def test_tax_excel_snapshot_and_replay(client, auth, app, worker):
    tax = create_tax(client, auth)
    row, headers, body = create_export(
        client, auth, selection={"mode": "ids", "ids": [tax["id"]]}, group_by="month"
    )
    assert row["record_count"] == 1 and row["pending_review_count"] == 0
    assert (
        client.post("/api/v1/exports", headers=headers, json=body).json()["data"]["id"] == row["id"]
    )
    create_tax(client, auth, "226")
    ready, content = finished(client, auth, worker, row)
    book = load_workbook(BytesIO(content), data_only=False)
    assert book.sheetnames == ["汇总", "拆税记录"]
    sheet = book["拆税记录"]
    assert (
        sheet.max_row == 2
        and sheet["B2"].value == 113
        and sheet["C2"].value == 100
        and sheet["D2"].value == 13
    )
    assert sheet["B2"].data_type == "n" and sheet.freeze_panes == "A2"
    assert book["汇总"]["C2"].value == 1
    assert (
        client.get("/api/v1/records?type=export", headers=auth).json()["data"]["items"][0][
            "resource_id"
        ]
        == row["id"]
    )
    assert client.get("/api/v1/capabilities").json()["data"]["export_enabled"]
    with app.state.session_factory() as db:
        assert db.scalar(select(func.count()).select_from(ExportSnapshot)) == 1
        assert db.scalar(select(func.count()).select_from(ExportFile)) == 1


def test_invoice_snapshot_pending_and_formula_safety(client, auth, app):
    batch, worker, _ = setup_batch(
        client, auth, app, callback=lambda req: normalize_response(payload(multi=True))
    )
    drain(worker)
    invoice = client.get("/api/v1/invoices", headers=auth).json()["data"]["items"][0]
    patched = client.patch(
        "/api/v1/invoices/" + invoice["id"],
        headers=auth,
        json={
            "expected_version": invoice["version"],
            "seller": {"name": ' =HYPERLINK("http://example","bad")'},
        },
    ).json()["data"]
    row, _, _ = create_export(
        client,
        auth,
        source="invoices",
        selection={"mode": "batch", "batch_id": batch["id"]},
        review_scope="all",
        group_by="seller",
    )
    client.patch(
        "/api/v1/invoices/" + invoice["id"],
        headers=auth,
        json={"expected_version": patched["version"], "seller": {"name": "后来修改"}},
    )
    ready, content = finished(client, auth, worker, row)
    assert ready["record_count"] == 2 and ready["pending_review_count"] == 2
    book = load_workbook(BytesIO(content), data_only=False)
    assert book.sheetnames == ["汇总", "票据明细", "待核对项目", "商品明细"]
    assert book["待核对项目"].max_row == 3 and book["商品明细"].max_row == 3
    texts = [row[5].value for row in book["票据明细"].iter_rows(min_row=2)]
    assert any(value.startswith("'=HYPERLINK") for value in texts) and "后来修改" not in texts
    assert all(cell.data_type != "f" for sheet in book for line in sheet for cell in line)
    assert book["票据明细"]["J2"].value == 113


def test_confirmed_default_empty_and_ownership(client, auth, app):
    batch, worker, _ = setup_batch(client, auth, app)
    drain(worker)
    body = {
        "format": "excel",
        "source_type": "invoices",
        "selection": {"mode": "batch", "batch_id": batch["id"]},
    }
    response = client.post(
        "/api/v1/exports", headers={**auth, "Idempotency-Key": "empty"}, json=body
    )
    assert response.status_code == 422 and response.json()["code"] == "EMPTY_EXPORT"
    with app.state.session_factory() as db:
        assert db.scalar(select(Export.id)) is None
    invoice = client.get("/api/v1/invoices", headers=auth).json()["data"]["items"][0]
    client.post(
        "/api/v1/invoices/" + invoice["id"] + "/confirm",
        headers=auth,
        json={"expected_version": invoice["version"]},
    )
    result = client.post(
        "/api/v1/exports", headers={**auth, "Idempotency-Key": "confirmed"}, json=body
    )
    assert result.status_code == 202 and result.json()["data"]["pending_review_count"] == 0
    other = client.post("/api/v1/auth/dev-login", json={"account": "bob"}).json()["data"]
    other_headers = {
        "Authorization": "Bearer " + other["access_token"],
        "Idempotency-Key": "foreign",
    }
    for selection in [
        {"mode": "ids", "ids": [invoice["id"]]},
        {"mode": "batch", "batch_id": batch["id"]},
    ]:
        assert (
            client.post(
                "/api/v1/exports",
                headers=other_headers,
                json={**body, "selection": selection, "review_scope": "all"},
            ).status_code
            == 404
        )
    id = result.json()["data"]["id"]
    for suffix in ["", "/download-url", "/share"]:
        assert (
            client.get("/api/v1/exports/" + id + suffix, headers=other_headers).status_code == 404
        )


@pytest.mark.parametrize(
    "patch",
    [
        {"selection": {"mode": "ids", "ids": []}},
        {"selection": {"mode": "ids", "ids": ["x", "x"]}},
        {"selection": {"mode": "batch", "batch_id": "x"}},
        {"group_by": "seller"},
        {"review_scope": "all"},
        {"selection": {"mode": "filter", "type_code": "vat_invoice"}},
        {"selection": {"mode": "filter", "date_from": "2026-10-10", "date_to": "2026-10-01"}},
        {"selection": {"mode": "filter", "sql": "SELECT *"}},
    ],
)
def test_invalid_export_requests(client, auth, patch):
    response = client.post(
        "/api/v1/exports",
        headers={**auth, "Idempotency-Key": "invalid"},
        json={
            "format": "excel",
            "source_type": "tax_calculations",
            "selection": {"mode": "filter"},
            **patch,
        },
    )
    assert response.status_code == 422


def test_share_replay_hashing_revocation_and_new_share(client, auth, app, worker):
    create_tax(client, auth)
    row, _, _ = create_export(client, auth)
    finished(client, auth, worker, row)
    endpoint = "/api/v1/exports/" + row["id"] + "/share"
    headers = {**auth, "Idempotency-Key": "share-once"}
    response = client.post(endpoint, headers=headers, json={"expires_in_days": 7})
    assert response.status_code == 201, response.text
    shared = response.json()["data"]
    token = shared["token"]
    assert (
        client.post(endpoint, headers=headers, json={"expires_in_days": 7}).json()["data"] == shared
    )
    assert (
        client.post(endpoint, headers={**auth, "Idempotency-Key": "another"}, json={}).json()[
            "code"
        ]
        == "SHARE_ALREADY_EXISTS"
    )
    info = client.get(endpoint, headers=auth).json()["data"]
    assert "token" not in info and "share_path" not in info
    public = client.get("/api/v1/shares/" + token).json()["data"]
    assert set(public) == {"filename", "format", "size_bytes", "expires_at", "description"}
    signed = client.get("/api/v1/shares/" + token + "/download-url").json()["data"]["url"]
    assert client.get(local_url(signed)).status_code == 200
    with app.state.session_factory() as db:
        record = db.scalar(
            select(IdempotencyRecord).where(
                IdempotencyRecord.operation == "share.create:" + row["id"]
            )
        )
        assert token not in json.dumps(record.response_data)
        assert db.scalar(select(ExportShare)).token_hash != token
    assert client.delete(endpoint, headers=auth).status_code == 200
    assert client.delete(endpoint, headers=auth).status_code == 200
    assert client.get("/api/v1/shares/" + token).status_code == 404
    assert client.get(local_url(signed)).status_code == 404
    another = client.post(endpoint, headers={**auth, "Idempotency-Key": "new"}, json={}).json()[
        "data"
    ]
    assert (
        another["token"] != token
        and client.get("/api/v1/shares/" + another["token"]).status_code == 200
    )
    assert client.get("/api/v1/shares/not-a-token").status_code == 404


def test_share_expiry_file_expiry_and_url_lifetime(client, auth, app, worker):
    create_tax(client, auth)
    row, _, _ = create_export(client, auth)
    finished(client, auth, worker, row)
    with app.state.session_factory.begin() as db:
        file = db.scalar(select(ExportFile).where(ExportFile.export_id == row["id"]))
        file.expires_at = now() + timedelta(seconds=60)
    shared = client.post(
        "/api/v1/exports/" + row["id"] + "/share",
        headers={**auth, "Idempotency-Key": "expires"},
        json={"expires_in_days": 30},
    ).json()["data"]
    signed = client.get("/api/v1/shares/" + shared["token"] + "/download-url").json()["data"]["url"]
    from urllib.parse import parse_qs, urlsplit

    claims = decode(
        app.state.settings, parse_qs(urlsplit(signed).query)["token"][0], "export_download"
    )
    assert claims["exp"] - claims["iat"] <= 60
    with app.state.session_factory.begin() as db:
        db.get(ExportShare, shared["id"]).expires_at = now() - timedelta(seconds=1)
    assert client.get("/api/v1/shares/" + shared["token"]).status_code == 410
    assert client.get(local_url(signed)).status_code == 410
    with app.state.session_factory.begin() as db:
        db.scalar(select(ExportFile).where(ExportFile.export_id == row["id"])).expires_at = (
            now() - timedelta(seconds=1)
        )
    assert (
        client.get("/api/v1/exports/" + row["id"], headers=auth).json()["data"]["file"]["status"]
        == "expired"
    )
    assert (
        client.get("/api/v1/exports/" + row["id"] + "/download-url", headers=auth).status_code
        == 410
    )


def test_cancel_and_expired_worker_fencing(client, auth, app, worker):
    create_tax(client, auth)
    row, _, _ = create_export(client, auth)
    old_id, old_token = worker.claim()
    with app.state.session_factory() as db:
        result = export_tasks.perform(
            db, db.get(Job, old_id), app.state.settings, app.state.storage
        )
    with app.state.session_factory.begin() as db:
        db.get(Job, old_id).lease_until = now() - timedelta(seconds=1)
    new_id, new_token = worker.claim()
    worker.finish(old_id, old_token, result, None)
    with app.state.session_factory() as db:
        assert db.scalar(select(ExportFile.id)) is None
    cancelled = client.post("/api/v1/exports/" + row["id"] + "/cancel", headers=auth, json={})
    assert cancelled.json()["data"]["status"] == "cancelling"
    worker.finish(new_id, new_token, result, None)
    assert (
        client.get("/api/v1/exports/" + row["id"], headers=auth).json()["data"]["status"]
        == "cancelled"
    )
    second, _, _ = create_export(client, auth)
    assert (
        client.post("/api/v1/exports/" + second["id"] + "/cancel", headers=auth, json={}).json()[
            "data"
        ]["status"]
        == "cancelled"
    )
    assert not worker.process_one()


def test_generation_retry_recovers(client, auth, app, worker, monkeypatch):
    create_tax(client, auth)
    row, _, _ = create_export(client, auth)
    publish = app.state.storage.publish
    calls = []

    def flaky(*args):
        calls.append(1)
        if len(calls) == 1:
            raise BusinessError("STORAGE_UNAVAILABLE", "temporary", 503, retryable=True)
        return publish(*args)

    monkeypatch.setattr(app.state.storage, "publish", flaky)
    assert worker.process_one()
    assert (
        client.get("/api/v1/exports/" + row["id"], headers=auth).json()["data"]["status"]
        == "queued"
    )
    with app.state.session_factory.begin() as db:
        db.scalar(select(Job).where(Job.resource_id == row["id"])).available_at = now() - timedelta(
            seconds=1
        )
    finished(client, auth, worker, row)
    assert len(calls) == 2


def test_currency_isolation_missing_values_and_large_excel_precision(client, auth, app):
    batch, worker, _ = setup_batch(
        client, auth, app, callback=lambda req: normalize_response(payload(multi=True))
    )
    drain(worker)
    rows = client.get("/api/v1/invoices", headers=auth).json()["data"]["items"]
    from app.db.models import Invoice
    from app.services.invoices import assign_data

    with app.state.session_factory.begin() as db:
        invoice = db.get(Invoice, rows[0]["id"])
        data = dict(invoice.data)
        data["currency"] = "UNKNOWN"
        data["gross_amount"] = "9999999999999999999999.99"
        data["tax_amount"] = "0.00"
        data["net_amount"] = None
        assign_data(invoice, data)
    row, _, _ = create_export(client, auth, source="invoices", review_scope="all")
    with app.state.session_factory() as db:
        values = summaries(db, db.get(Export, row["id"]))
        assert values[("全部", "CNY")]["gross_amount"] == Decimal("113")
        assert values[("全部", "UNKNOWN")]["gross_amount"] is None
    ready, content = finished(client, auth, worker, row)
    book = load_workbook(BytesIO(content))
    big = next(
        line for line in book["票据明细"].iter_rows(min_row=2) if line[12].value == "UNKNOWN"
    )
    assert big[9].value == "9999999999999999999999.99" and big[9].data_type == "s"
    assert big[10].value is None and big[11].value == 0


def test_pdf_chinese_multipage_and_escaped_text(client, auth, app):
    batch, worker, _ = setup_batch(client, auth, app, pages=35)
    drain(worker)
    row, _, _ = create_export(
        client, auth, format="pdf", source="invoices", review_scope="all", group_by="month"
    )
    ready, content = finished(client, auth, worker, row)
    pdf = PdfReader(BytesIO(content))
    text = "".join(page.extract_text() for page in pdf.pages)
    assert len(pdf.pages) >= 2 and "金融工具箱" in text and "待核对" in text
    assert "供应商" in text and "113.00" in text


def test_export_tool_matches_page_idempotency(client, auth, app):
    from app.agent.tools import ToolContext, build_registry

    create_tax(client, auth)
    user = client.get("/api/v1/users/me", headers=auth).json()["data"]["id"]
    args = {"format": "excel", "source_type": "tax_calculations", "selection": {"mode": "filter"}}
    with app.state.session_factory.begin() as db:
        ctx = ToolContext(db, user, "export-step")
        registry = build_registry()
        first = registry.execute("create_export", args, ctx)
        assert registry.execute("create_export", args, ctx) == first
        assert (
            registry.execute("get_export", {"resource_id": first["id"]}, ctx)["id"] == first["id"]
        )
        with pytest.raises(BusinessError):
            registry.execute("create_share", {"resource_id": first["id"]}, ctx)


def test_excel_sheet_limit_splits_without_losing_records(client, auth, app, worker, monkeypatch):
    from app.services import export_render

    for _ in range(15):
        create_tax(client, auth)
    monkeypatch.setattr(export_render, "MAX_SHEET_ROWS", 8)
    row, _, _ = create_export(client, auth)
    _, content = finished(client, auth, worker, row)
    book = load_workbook(BytesIO(content))
    sheets = [sheet for sheet in book if sheet.title.startswith("拆税记录")]
    assert len(sheets) == 3
    assert sum(sheet.max_row - 1 for sheet in sheets) == 15


@pytest.mark.parametrize("text", ["=SUM(1,2)", " +CMD", "\t@HYPERLINK()", "-TEXT"])
def test_text_injection_guard(text):
    from app.services.export_render import safe_text

    assert safe_text(text).startswith("'")


def test_cos_generated_export_and_signed_download(client, auth, app):
    from app.jobs.worker import Worker
    from tests.test_cos import setup_cos

    storage, cos = setup_cos(app)
    uploads = []

    def upload_file(**kw):
        uploads.append(kw)
        cos.objects[kw["Key"]] = __import__("pathlib").Path(kw["LocalFilePath"]).read_bytes()
        return {}

    cos.upload_file = upload_file
    create_tax(client, auth)
    row, _, _ = create_export(client, auth)
    worker = Worker(app.state.settings, app.state.session_factory, storage)
    drain(worker)
    ready = client.get("/api/v1/exports/" + row["id"], headers=auth).json()["data"]
    assert ready["status"] == "succeeded" and len(uploads) == 1
    assert (
        uploads[0]["ContentType"]
        == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    assert cos.objects[uploads[0]["Key"]].startswith(b"PK")
    shared = client.post(
        "/api/v1/exports/" + row["id"] + "/share",
        headers={**auth, "Idempotency-Key": "cos-share"},
        json={},
    ).json()["data"]
    link = client.get("/api/v1/shares/" + shared["token"] + "/download-url")
    assert link.status_code == 200 and cos.signed[-1]["Method"] == "GET"
    assert cos.signed[-1]["Key"] == uploads[0]["Key"] and cos.signed[-1]["Expired"] <= 300


def test_pdf_preserves_long_escaped_seller_text(client, auth, app):
    value = payload()
    value["MixedInvoiceItems"][0]["SingleInvoiceInfos"]["VatSpecialInvoice"]["Seller"] = (
        "特殊<&>字符" + ("长文本销售方" * 130)
    )
    batch, worker, _ = setup_batch(
        client, auth, app, callback=lambda req: normalize_response(value)
    )
    drain(worker)
    row, _, _ = create_export(
        client, auth, format="pdf", source="invoices", review_scope="all", group_by="seller"
    )
    _, content = finished(client, auth, worker, row)
    text = "".join(page.extract_text() for page in PdfReader(BytesIO(content)).pages)
    assert "特殊<&>字符" in text and text.count("长") == 260
