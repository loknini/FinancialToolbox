import math
import unicodedata
from datetime import date, datetime
from decimal import Decimal, localcontext
from xml.sax.saxutils import escape
from zoneinfo import ZoneInfo

from openpyxl import Workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.styles import Alignment, Font, PatternFill
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import ParagraphStyle
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.cidfonts import UnicodeCIDFont
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import LongTable, Paragraph, SimpleDocTemplate, Spacer, TableStyle
from sqlalchemy import select

from app.core.security import timestamp
from app.db.models import ExportSnapshot

MAX_SHEET_ROWS = 1_048_576

MONEY_FIELDS = ("gross_amount", "net_amount", "tax_amount")


def records(db, export):
    return db.scalars(
        select(ExportSnapshot.data)
        .where(ExportSnapshot.export_id == export.id)
        .order_by(ExportSnapshot.position)
        .execution_options(yield_per=500)
    )


def grouping(row, mode):
    if mode == "seller":
        return row.get("seller", {}).get("name") or "销售方待核对"
    if mode == "month":
        if "invoice_date" in row:
            return row.get("invoice_date", "")[:7] if row.get("invoice_date") else "日期待核对"
        return (
            datetime.fromisoformat(row["created_at"].replace("Z", "+00:00"))
            .astimezone(ZoneInfo("Asia/Shanghai"))
            .strftime("%Y-%m")
        )
    return "全部"


def summaries(db, export):
    result = {}
    with localcontext() as context:
        context.prec = 60
        for row in records(db, export):
            key = (grouping(row, export.request_data["group_by"]), row.get("currency", "CNY"))
            if key not in result:
                result[key] = {
                    "count": 0,
                    "pending": 0,
                    **{k: Decimal(0) for k in MONEY_FIELDS},
                    **{k + "_missing": 0 for k in MONEY_FIELDS},
                }
            value = result[key]
            value["count"] += 1
            value["pending"] += int(row.get("review_status") == "pending_review")
            for field in MONEY_FIELDS:
                if row.get(field) is None:
                    value[field + "_missing"] += 1
                else:
                    value[field] += Decimal(row[field])
        for (_, currency), value in result.items():
            for field in MONEY_FIELDS:
                if currency == "UNKNOWN" or value[field + "_missing"] == value["count"]:
                    value[field] = None
    return result


def safe_text(value):
    if value is None:
        return None
    text = str(value)
    text = "".join(char for char in text if ord(char) >= 32 or char in "\t\n\r")
    if text.lstrip().startswith(("=", "+", "-", "@")):
        text = "'" + text
    return text


def amount(value):
    if value is None:
        return None
    number = Decimal(value)
    # Excel has 15 significant digits. Keep larger values as exact text.
    return format(number, ".2f") if len(number.normalize().as_tuple().digits) > 15 else number


def excel(db, export, path):
    book = Workbook(write_only=True)
    sheet_meta = {}
    current_sheets = {}
    row_counts = {}
    header_fill = PatternFill("solid", fgColor="173F5F")

    def sheet(name, headers, widths):
        ws = book.create_sheet(name)
        sheet_meta[ws.title] = (headers, widths)
        row_counts[ws.title] = 1
        current_sheets[ws.title] = ws
        ws.freeze_panes = "A2"
        for index, width in enumerate(widths, 1):
            from openpyxl.utils import get_column_letter

            ws.column_dimensions[get_column_letter(index)].width = width
        cells = []
        for text in headers:
            cell = WriteOnlyCell(ws, value=text)
            cell.font = Font(name="微软雅黑", bold=True, color="FFFFFF")
            cell.fill = header_fill
            cell.alignment = Alignment(wrap_text=True, vertical="center")
            cells.append(cell)
        ws.append(cells)
        ws.row_dimensions[1].height = 30
        return ws

    def append(
        ws,
        values,
        money_columns=(),
        pending=False,
        decimal_columns=(),
        rate_columns=(),
        date_columns=(),
        time_columns=(),
        merge_note=False,
    ):
        root = ws.title
        ws = current_sheets[root]
        headers, widths = sheet_meta[root]
        if row_counts[ws.title] >= MAX_SHEET_ROWS:
            suffix = str(1 + sum(title.startswith(root) for title in book.sheetnames))
            ws = sheet(root[: 31 - len(suffix)] + suffix, headers, widths)
            current_sheets[root] = ws
        row_counts[ws.title] += 1
        height = 1
        for index, value in enumerate(values):
            text = "" if value is None else str(value)
            height = max(
                height,
                sum(
                    max(
                        1,
                        math.ceil(
                            sum(
                                2 if unicodedata.east_asian_width(c) in ("W", "F") else 1
                                for c in part
                            )
                            / max(8, widths[index] - 2)
                        ),
                    )
                    for part in text.split("\n")
                ),
            )
        ws.row_dimensions[row_counts[ws.title]].height = min(409, max(20, height * 14 + 4))
        if merge_note:
            row_number = row_counts[ws.title]
            ws.merged_cells.add(f"B{row_number}:J{row_number}")
            ws.row_dimensions[row_number].height = 30
        cells = []
        for index, value in enumerate(values):
            item = (
                amount(value)
                if index in money_columns
                else value
                if isinstance(value, int)
                else safe_text(value)
            )
            if value is not None and index in (*decimal_columns, *rate_columns):
                number = Decimal(value)
                item = (
                    format(number, "f")
                    if len(number.normalize().as_tuple().digits) > 15
                    else number
                )
            if value is not None and index in date_columns:
                item = date.fromisoformat(value)
            if value is not None and index in time_columns:
                item = (
                    datetime.fromisoformat(value.replace("Z", "+00:00"))
                    .astimezone(ZoneInfo("Asia/Shanghai"))
                    .replace(tzinfo=None)
                )
            cell = WriteOnlyCell(ws, value=item)
            if isinstance(item, str):
                cell.data_type = "s"
            cell.font = Font(name="微软雅黑", size=10)
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            if index in money_columns:
                cell.number_format = "#,##0.00;[Red]-#,##0.00"
            if index in decimal_columns:
                cell.number_format = "General"
            if index in rate_columns:
                cell.number_format = "0.00%"
            if index in date_columns:
                cell.number_format = "yyyy-mm-dd"
            if index in time_columns:
                cell.number_format = "yyyy-mm-dd hh:mm:ss"
            if index in (*money_columns, *decimal_columns, *rate_columns) and isinstance(item, str):
                cell.number_format = "@"
            if pending:
                cell.fill = PatternFill("solid", fgColor="FFF2CC")
            cells.append(cell)
        ws.append(cells)

    summary = sheet(
        "汇总",
        [
            "分组",
            "币种",
            "记录数",
            "待核对数",
            "含税额已知部分",
            "不含税额已知部分",
            "税额已知部分",
            "缺含税额数",
            "缺不含税额数",
            "缺税额数",
        ],
        [28, 12, 12, 12, 23, 23, 23, 16, 18, 16],
    )
    for (group, currency), value in sorted(summaries(db, export).items()):
        append(
            summary,
            [
                group,
                currency,
                value["count"],
                value["pending"],
                *[value[k] for k in MONEY_FIELDS],
                *[value[k + "_missing"] for k in MONEY_FIELDS],
            ],
            (4, 5, 6),
        )
    append(summary, ["快照时间", timestamp(export.snapshot_at)], merge_note=True)
    append(
        summary,
        ["说明", "金额为快照中的已知部分；缺失值为空，不补零。未知币种不合计。"],
        merge_note=True,
    )
    append(
        summary,
        ["精度", "超过 Excel 15 位有效数字的金额按精确文本保存，请勿转为数值。"],
        merge_note=True,
    )
    append(
        summary,
        ["核对范围", export.request_data.get("review_scope") or "拆税记录"],
        merge_note=True,
    )
    if export.source_type == "tax_calculations":
        ws = sheet(
            "拆税记录",
            [
                "记录ID",
                "含税金额",
                "不含税金额",
                "税额",
                "税率",
                "币种",
                "来源票据ID",
                "创建时间（北京时间）",
                "规则版本",
            ],
            [40, 23, 23, 23, 12, 12, 40, 30, 24],
        )
        for row in records(db, export):
            append(
                ws,
                [
                    row["id"],
                    row["gross_amount"],
                    row["net_amount"],
                    row["tax_amount"],
                    row["tax_rate"],
                    "CNY",
                    row.get("source_invoice_id"),
                    row["created_at"],
                    row["rule_version"],
                ],
                (1, 2, 3),
                rate_columns=(4,),
                time_columns=(7,),
            )
    else:
        headers = [
            "票据ID",
            "票种",
            "票据号码",
            "代码",
            "票面日期",
            "销售方",
            "销售方税号",
            "购买方",
            "购买方税号",
            "含税金额",
            "不含税金额",
            "税额",
            "币种",
            "核对状态",
            "核对问题",
            "来源文件ID",
            "来源页码",
        ]
        widths = [40, 24, 28, 22, 16, 38, 28, 38, 28, 23, 23, 23, 12, 16, 56, 40, 12]
        ws = sheet("票据明细", headers, widths)
        pending_ws = sheet("待核对项目", headers, widths)
        line_ws = sheet(
            "商品明细",
            [
                "票据ID",
                "行号",
                "商品/服务",
                "数量",
                "单价",
                "不含税金额",
                "税率",
                "税收标识",
                "税额",
                "币种",
            ],
            [40, 12, 48, 20, 23, 23, 14, 20, 23, 12],
        )
        for row in records(db, export):
            issues = "；".join(i.get("message", "") for i in row.get("issues", []))
            values = [
                row["id"],
                row["type_label"],
                row.get("invoice_number"),
                row.get("invoice_code"),
                row.get("invoice_date"),
                row.get("seller", {}).get("name"),
                row.get("seller", {}).get("tax_id"),
                row.get("buyer", {}).get("name"),
                row.get("buyer", {}).get("tax_id"),
                *[row.get(k) for k in MONEY_FIELDS],
                row.get("currency", "CNY"),
                "已确认" if row["review_status"] == "confirmed" else "待核对",
                issues,
                row.get("source", {}).get("file_id"),
                row.get("source", {}).get("page_from"),
            ]
            pending = row["review_status"] == "pending_review"
            append(ws, values, (9, 10, 11), pending, date_columns=(4,))
            if pending:
                append(pending_ws, values, (9, 10, 11), True, date_columns=(4,))
            for line in row.get("lines", []):
                append(
                    line_ws,
                    [
                        row["id"],
                        line.get("line_number"),
                        line.get("name"),
                        line.get("quantity"),
                        line.get("unit_price"),
                        line.get("net_amount"),
                        line.get("tax_rate"),
                        line.get("tax_treatment"),
                        line.get("tax_amount"),
                        row.get("currency", "CNY"),
                    ],
                    (5, 8),
                    pending,
                    decimal_columns=(3, 4),
                    rate_columns=(6,),
                )
    book.save(path)


def pdf(db, export, path, settings):
    font = "STSong-Light"
    if settings.pdf_font_path:
        font = "ExportChinese"
        pdfmetrics.registerFont(TTFont(font, str(settings.pdf_font_path)))
    elif font not in pdfmetrics.getRegisteredFontNames():
        pdfmetrics.registerFont(UnicodeCIDFont(font))
    style = ParagraphStyle(
        "body",
        fontName=font,
        fontSize=9,
        leading=13,
        wordWrap="CJK",
        textColor=colors.HexColor("#16324f"),
    )
    title = ParagraphStyle("title", parent=style, fontSize=20, leading=26, spaceAfter=12)

    def p(value):
        return Paragraph(escape(str(value if value is not None else "待核对")), style)

    class AlignedTable(LongTable):
        def split(self, available_width, available_height):
            parts = super().split(available_width, available_height)

            for part in parts:
                part.hAlign = "LEFT"
            return parts

    def table(rows, widths):
        grid = AlignedTable(
            [[p(value) for value in row] for row in rows],
            colWidths=widths,
            repeatRows=1,
            splitInRow=1,
            hAlign="LEFT",
        )
        grid.setStyle(
            TableStyle(
                [
                    ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#e4edf3")),
                    ("GRID", (0, 0), (-1, -1), 0.3, colors.HexColor("#c5d4dd")),
                    ("VALIGN", (0, 0), (-1, -1), "TOP"),
                    ("TOPPADDING", (0, 0), (-1, -1), 6),
                    ("BOTTOMPADDING", (0, 0), (-1, -1), 6),
                ]
            )
        )
        return grid

    story = [
        Paragraph("金融工具箱 · 导出报告", title),
        p("快照时间：" + timestamp(export.snapshot_at)),
        p(f"记录数：{export.record_count}  |  待核对：{export.pending_review_count}"),
        p("金额汇总仅含已知部分；未知币种不合计。空值需核对，不能视为零。"),
        Spacer(1, 14),
    ]
    summary_rows = [
        [
            "分组",
            "币种",
            "记录/待核对",
            "含税额已知部分",
            "不含税额已知部分",
            "税额已知部分",
            "缺失数（含/净/税）",
        ]
    ]
    for (group, currency), value in sorted(summaries(db, export).items()):
        summary_rows.append(
            [
                group,
                currency,
                f"{value['count']} / {value['pending']}",
                *[
                    format(value[k], ".2f") if value[k] is not None else "未合计"
                    for k in MONEY_FIELDS
                ],
                "/".join(str(value[k + "_missing"]) for k in MONEY_FIELDS),
            ]
        )
    story.extend([table(summary_rows, [130, 70, 75, 125, 125, 115, 70]), Spacer(1, 16)])
    if export.source_type == "invoices":
        rows = [
            [
                "号码/票种",
                "票面日期",
                "销售方",
                "含税金额",
                "不含税金额",
                "税额",
                "币种",
                "核对状态/问题",
            ]
        ]
        for row in records(db, export):
            issues = "；".join(i.get("message", "") for i in row.get("issues", []))
            rows.append(
                [
                    str(row.get("invoice_number") or "号码待核对") + " / " + row["type_label"],
                    row.get("invoice_date"),
                    row.get("seller", {}).get("name"),
                    *[row.get(k) for k in MONEY_FIELDS],
                    row.get("currency", "CNY"),
                    ("已确认" if row["review_status"] == "confirmed" else "待核对")
                    + ("；" + issues if issues else ""),
                ]
            )
        detail_table = table(rows, [110, 65, 145, 85, 85, 80, 70, 70])
        story.append(detail_table)
    else:
        rows = [["创建时间", "含税金额", "不含税金额", "税额", "税率", "规则版本"]]
        for row in records(db, export):
            rows.append(
                [
                    row["created_at"],
                    *[row[k] for k in MONEY_FIELDS],
                    row["tax_rate"],
                    row["rule_version"],
                ]
            )
        detail_table = table(rows, [175, 135, 135, 135, 55, 75])
        story.append(detail_table)

    def footer(canvas, doc):
        canvas.setFont(font, 8)
        canvas.setFillColor(colors.HexColor("#5c7180"))
        canvas.drawString(32, 20, "金融工具箱 | 此文件为汇总报告，非原始票据或真伪查验结果")
        canvas.drawRightString(810, 20, "第 " + str(doc.page) + " 页")

    SimpleDocTemplate(
        str(path),
        pagesize=landscape(A4),
        leftMargin=32,
        rightMargin=32,
        topMargin=30,
        bottomMargin=40,
        title="金融工具箱导出报告",
        author="金融工具箱",
    ).build(story, onFirstPage=footer, onLaterPages=footer)

    # ReportLab may emit a final, height-only row continuation. Inspect actual
    # painted spans, never Paragraph.text, whose split fragments can be empty.
    import pymupdf

    header_bottom = 30 + detail_table._rowHeights[0] - 2
    trimmed = path.with_name(path.stem + "-trimmed.pdf")
    with pymupdf.open(path) as document:
        removed = False
        while len(document) > 1:
            page = document[-1]
            body = [
                span
                for block in page.get_text("dict")["blocks"]
                if "lines" in block
                for line in block["lines"]
                for span in line["spans"]
                if span["text"].strip() and header_bottom <= span["bbox"][1] < page.rect.height - 45
            ]
            if body:
                break
            document.delete_page(len(document) - 1)
            removed = True
        if removed:
            document.save(trimmed)
    if trimmed.exists():
        trimmed.replace(path)
